import os
from docx import Document
import re

replacements = [
    # Supervisor name
    (r'Nguyễn Thị Hoàng Yến', 'Phan Duy Hùng'),
    # Old project name still snagged in some files
    (r'"Time-Aware Explainable ML for Early At-Risk Student Prediction on OULAD"', '"EduGuard: AI-Powered Early Warning and Actionable Recourse System for Students"'),
    (r'Time-Aware Explainable ML for Early At-Risk Student Prediction on OULAD', 'EduGuard: AI-Powered Early Warning and Actionable Recourse System for Students'),
    # Old team list in PROJECT_STATUS.md
    (r'6 thành viên \(Phúc, Đức, Khoa, Bình, Huy Anh, Vinh\)', '4 thành viên (Vinh Lê, Ngọc Sang, Văn Sơn, Huy Anh)'),
    (r'6 thành viên \(Vinh Lê, Ngọc Sang, Văn Sơn, Huy Anh, Huy Anh, Vinh\)', '4 thành viên (Vinh Lê, Ngọc Sang, Văn Sơn, Huy Anh)'),
    # Fix README which still says Group 1
    (r'Group 1, DSP391m', 'Group 5, DSP391m'),
    (r'Nhóm 1, DSP391m', 'Nhóm 5, DSP391m'),
]

dirs_to_process = ['.']  # Root level files
extra_files = [
    'README.md',
    'PROJECT_STATUS.md',
    'ARCHITECTURE.md',
    'SETUP_VI.md',
]

def replace_in_text(content):
    for old, new in replacements:
        content = re.sub(old, new, content)
    return content

def process_text_file(filepath):
    try:
        with open(filepath, 'r', encoding='utf-8') as f:
            content = f.read()
        new_content = replace_in_text(content)
        if new_content != content:
            with open(filepath, 'w', encoding='utf-8') as f:
                f.write(new_content)
            print(f"Updated: {filepath}")
    except Exception as e:
        pass

def process_docx_file(filepath):
    try:
        doc = Document(filepath)
        changed = False
        for p in doc.paragraphs:
            new_text = replace_in_text(p.text)
            if new_text != p.text:
                p.text = new_text
                changed = True
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        new_text = replace_in_text(p.text)
                        if new_text != p.text:
                            p.text = new_text
                            changed = True
        if changed:
            doc.save(filepath)
            print(f"Updated docx: {filepath}")
    except Exception as e:
        pass

# Process root-level files
for fname in extra_files:
    fpath = os.path.join(os.getcwd(), fname)
    if os.path.exists(fpath):
        process_text_file(fpath)

# Process all subdirs
for dirn in ['docs', 'reports', 'tools', 'src', 'notebooks']:
    dirpath = os.path.join(os.getcwd(), dirn)
    if not os.path.exists(dirpath):
        continue
    for root, dirs, files in os.walk(dirpath):
        for file in files:
            fp = os.path.join(root, file)
            ext = os.path.splitext(file)[1].lower()
            if ext in ['.md', '.py', '.txt', '.tex', '.csv']:
                process_text_file(fp)
            elif ext == '.docx':
                process_docx_file(fp)

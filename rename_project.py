import os
from docx import Document

# Define the strings to replace
replacements = {
    "Time-Aware Explainable Machine Learning for Early At-Risk Student Prediction on OULAD": "EduGuard: AI-Powered Early Warning and Actionable Recourse System for Students",
    "Time-Aware Explainable Machine Learning for Early\\n                  At-Risk Student Prediction on OULAD": "EduGuard: AI-Powered Early Warning and Actionable Recourse System for Students",
    "Time-Aware Explainable Machine Learning for Early\n                  At-Risk Student Prediction on OULAD": "EduGuard: AI-Powered Early Warning and Actionable Recourse System for Students",
    "Phát hiện sớm sinh viên có nguy cơ học tập kém bằng học máy có khả năng giải thích": "EduGuard: Hệ thống Cảnh báo Sớm và Tư vấn Học tập tự động dựa trên AI",
    "Sơn · Khoa · An · Đức · Phúc · Bình": "Vinh Lê · Ngọc Sang · Văn Sơn · Huy Anh",
    "Khoa, Bình, Đức, Phúc, Sơn, An": "Vinh Lê, Ngọc Sang, Bùi Văn Sơn, Lê Huy Anh",
    # Single name replacements (careful with short names, but okay for this specific repo)
    "Phúc": "Vinh Lê",
    "Đức": "Ngọc Sang",
    "Khoa": "Văn Sơn",
    "Bình": "Huy Anh",
    "Sơn": "Vinh",
    " An ": " Sang "
}

# Directories to process
dirs_to_process = ['docs', 'reports', 'tools']

def replace_in_text_file(filepath):
    try:
        with open(filepath, 'r', encoding='utf-8') as f:
            content = f.read()
            
        original_content = content
        for old, new in replacements.items():
            content = content.replace(old, new)
            
        if content != original_content:
            with open(filepath, 'w', encoding='utf-8') as f:
                f.write(content)
            print(f"Updated text file: {filepath}")
    except Exception as e:
        pass # Ignore encoding errors for non-utf8 files

def replace_in_docx_file(filepath):
    try:
        doc = Document(filepath)
        changed = False
        
        # Replace in paragraphs
        for p in doc.paragraphs:
            for old, new in replacements.items():
                if old in p.text:
                    p.text = p.text.replace(old, new)
                    changed = True
                    
        # Replace in tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for old, new in replacements.items():
                            if old in p.text:
                                p.text = p.text.replace(old, new)
                                changed = True
                                
        if changed:
            doc.save(filepath)
            print(f"Updated docx file: {filepath}")
    except Exception as e:
        print(f"Error processing docx file {filepath}: {e}")

# Crawl through directories
for dir_name in dirs_to_process:
    dir_path = os.path.join(os.getcwd(), dir_name)
    if not os.path.exists(dir_path):
        continue
        
    for root, dirs, files in os.walk(dir_path):
        for file in files:
            filepath = os.path.join(root, file)
            ext = os.path.splitext(file)[1].lower()
            
            if ext in ['.md', '.py', '.txt', '.tex', '.csv']:
                replace_in_text_file(filepath)
            elif ext == '.docx':
                replace_in_docx_file(filepath)

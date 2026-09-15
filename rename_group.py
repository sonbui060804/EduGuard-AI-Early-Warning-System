import os
from docx import Document
import re

replacements = [
    (r'\bGroup 1\b', 'Group 5'),
    (r'\bNhóm 1\b', 'Nhóm 5'),
    (r'\bgroup1\b', 'group5'),
    (r'\bNhom1\b', 'Nhom5')
]

dirs_to_process = ['docs', 'reports', 'tools']

def replace_in_text_file(filepath):
    try:
        with open(filepath, 'r', encoding='utf-8') as f:
            content = f.read()
            
        original_content = content
        for old, new in replacements:
            content = re.sub(old, new, content, flags=re.IGNORECASE)
            
        if content != original_content:
            with open(filepath, 'w', encoding='utf-8') as f:
                f.write(content)
            print(f"Updated text file: {filepath}")
    except Exception as e:
        pass

def replace_in_docx_file(filepath):
    try:
        doc = Document(filepath)
        changed = False
        
        # Replace in paragraphs
        for p in doc.paragraphs:
            for old, new in replacements:
                if re.search(old, p.text, re.IGNORECASE):
                    p.text = re.sub(old, new, p.text, flags=re.IGNORECASE)
                    changed = True
                    
        # Replace in tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        for old, new in replacements:
                            if re.search(old, p.text, re.IGNORECASE):
                                p.text = re.sub(old, new, p.text, flags=re.IGNORECASE)
                                changed = True
                                
        if changed:
            doc.save(filepath)
            print(f"Updated docx file: {filepath}")
    except Exception as e:
        pass

for dir_name in dirs_to_process:
    dir_path = os.path.join(os.getcwd(), dir_name)
    if not os.path.exists(dir_path):
        continue
        
    for root, dirs, files in os.walk(dir_path):
        for file in files:
            filepath = os.path.join(root, file)
            ext = os.path.splitext(file)[1].lower()
            
            if ext in ['.md', '.py', '.txt', '.tex', '.csv']:
                replace_in_text_file(filepath)
            elif ext == '.docx':
                replace_in_docx_file(filepath)
